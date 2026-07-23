"""Typed business-confirmation flow tests."""

from __future__ import annotations

import copy
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID

import pandas as pd
import pytest
from httpx import AsyncClient
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.db.tables import (
    AnalysisRun,
    Conversation,
    Message,
    PreflightReportRecord,
    Project,
    ProjectDataSource,
    SanitationRecipeRecord,
    SemanticEntry,
)
from app.models import SSEEvent
from app.services.analysis_checkpoint import stable_payload_hash
from app.services.correction_targets import discover_report_correction_targets
from app.services.data_preflight import fingerprint_file
from app.services.execution import ExecutionService
from app.services.semantic_revisions import append_semantic_revision


async def _seed_confirmation_run(
    db: AsyncSession,
) -> tuple[Project, Conversation, AnalysisRun]:
    project = Project(name="确认流项目")
    conversation = Conversation(title="收入调查", status="completed")
    db.add_all([project, conversation])
    await db.flush()
    run = AnalysisRun(
        project_id=project.id,
        conversation_id=conversation.id,
        query="比较各商品品类与门店的订单量",
        state="waiting_confirmation",
        stage="waiting_confirmation",
        report={
            "status": "waiting_confirmation",
            "title": "需要确认收入口径",
            "summary": "退款口径会改变结论。",
            "confirmation": {
                "key": "revenue_refund_policy",
                "question": "计算收入时，退款订单需要扣除吗？",
                "options": ["扣除退款", "保留退款订单"],
                "reason": "两种口径会改变收入和门店排名。",
            },
        },
        checkpoint={"resumable": False, "reason": "waiting_confirmation"},
    )
    db.add(run)
    await db.commit()
    await db.refresh(run)
    return project, conversation, run


def _command(run: AnalysisRun, **overrides: Any) -> dict[str, str]:
    payload = {
        "analysis_run_id": str(run.id),
        "key": "revenue_refund_policy",
        "selected_option": "扣除退款",
    }
    payload.update(overrides)
    return payload


async def _seed_sheet_confirmation_run(
    client: AsyncClient,
    db: AsyncSession,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[Project, Conversation, AnalysisRun, ProjectDataSource, dict[str, Any]]:
    monkeypatch.setattr(settings, "WORKSPACE_ROOT", tmp_path / "projects")
    project = Project(name="工作表确认项目")
    db.add(project)
    await db.commit()

    workbook = Workbook()
    orders = workbook.active
    orders.title = "订单明细"
    orders.append(["订单号", "门店ID", "商品", "实付"])
    orders.append(["O-1", "S-1", "A类", 32])
    orders.append(["O-2", "S-2", "B类", 28])
    orders.append(["O-3", "S-1", "A类", 30])
    refunds = workbook.create_sheet("退款明细")
    refunds.append(["退款单号", "渠道", "原因"])
    refunds.append(["R-1", "小程序", "重复付款"])
    refunds.append(["R-2", "外卖", "缺货"])
    refunds.append(["R-3", "小程序", "用户取消"])
    payload = BytesIO()
    workbook.save(payload)

    uploaded = await client.post(
        f"/api/v1/projects/{project.id}/sources/files",
        files={
            "file": (
                "two-sheets.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 200, uploaded.text
    source_id = uploaded.json()["data"]["id"]
    preflight = await client.post(f"/api/v1/projects/{project.id}/sources/{source_id}/preflight")
    assert preflight.status_code == 200, preflight.text
    report = preflight.json()["data"]
    question = next(
        item for item in report["ambiguities"] if item["key"].startswith("excel_sheet_selection:")
    )
    assert question["key"] == f"excel_sheet_selection:{source_id}"
    assert set(question["options"]) == {"订单明细", "退款明细"}
    assert report["source_snapshot"]["reader"]["selected_sheet"] == "订单明细"

    conversation = Conversation(title="工作表调查", status="completed")
    db.add(conversation)
    await db.flush()
    run = AnalysisRun(
        project_id=project.id,
        conversation_id=conversation.id,
        query="分析退款明细",
        state="waiting_confirmation",
        stage="waiting_confirmation",
        report={
            "status": "waiting_confirmation",
            "title": "需要确认工作表",
            "summary": "工作簿包含多个数据表。",
            "confirmation": question,
        },
        checkpoint={"resumable": False, "reason": "waiting_confirmation"},
    )
    db.add(run)
    await db.commit()
    source = await db.get(ProjectDataSource, UUID(source_id))
    assert source is not None
    return project, conversation, run, source, question


@pytest.mark.asyncio
async def test_typed_confirmation_persists_knowledge_message_and_resolves_preflight(
    client: AsyncClient,
    db_session: AsyncSession,
):
    project, conversation, run = await _seed_confirmation_run(db_session)
    source = ProjectDataSource(
        project_id=project.id,
        kind="file",
        name="orders.xlsx",
        format="xlsx",
        status="needs_confirmation",
        profile_data={
            "ambiguities": [
                {
                    "key": "revenue_refund_policy",
                    "question": "退款是否计入收入？",
                    "reason": "会影响收入",
                    "options": ["扣除退款", "保留退款订单"],
                }
            ],
            "summary": "数据已整理，有 1 个业务口径需要确认",
        },
    )
    db_session.add(source)
    await db_session.flush()
    preflight = PreflightReportRecord(
        project_id=project.id,
        data_source_id=source.id,
        status="needs_confirmation",
        summary="数据已整理，有 1 个业务口径需要确认",
        ambiguities=list(source.profile_data["ambiguities"]),
    )
    db_session.add(preflight)
    await db_session.commit()

    response = await client.post("/api/v1/chat/confirm", json=_command(run))

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["success"] is True
    assert body["data"]["analysis_run_id"] == str(run.id)
    assert body["data"]["resume_run_id"] == str(run.id)
    assert body["data"]["ready_to_continue"] is True
    await db_session.refresh(run)
    assert run.state == "waiting_confirmation"
    assert run.stage == "confirmation_received"
    assert run.checkpoint["confirmation_receipt_status"] == "pending"
    assert run.checkpoint["confirmation_receipt"]["task_query"] == run.query

    entry = (
        await db_session.execute(
            select(SemanticEntry).where(
                SemanticEntry.project_id == project.id,
                SemanticEntry.key == "revenue_refund_policy",
            )
        )
    ).scalar_one()
    assert entry.value == "扣除退款"
    assert entry.state == "confirmed"
    assert entry.source == "user"
    receipt = run.checkpoint["confirmation_receipt"]
    assert receipt["semantic_entry_id"] == str(entry.id)
    assert receipt["active_revision_id"] == str(entry.active_revision_id)
    assert receipt["definition_hash"] == stable_payload_hash(entry.definition)
    assert receipt["value_hash"] == stable_payload_hash(entry.value)
    message = (
        await db_session.execute(select(Message).where(Message.conversation_id == conversation.id))
    ).scalar_one()
    assert message.content == "扣除退款"
    assert message.extra_data["analysis_run_id"] == str(run.id)
    await db_session.refresh(preflight)
    await db_session.refresh(source)
    assert preflight.status == "ready"
    assert preflight.ambiguities == []
    assert source.status == "ready"
    assert source.profile_data["ambiguities"] == []


@pytest.mark.asyncio
async def test_typed_confirmation_is_idempotent_for_same_answer_and_rejects_another(
    client: AsyncClient,
    db_session: AsyncSession,
):
    project, conversation, run = await _seed_confirmation_run(db_session)

    first = await client.post("/api/v1/chat/confirm", json=_command(run))
    repeated = await client.post("/api/v1/chat/confirm", json=_command(run))
    conflicting = await client.post(
        "/api/v1/chat/confirm",
        json=_command(run, selected_option="保留退款订单"),
    )

    assert first.status_code == 200, first.text
    assert repeated.status_code == 200, repeated.text
    assert conflicting.status_code == 409, conflicting.text
    entry = (
        await db_session.execute(
            select(SemanticEntry).where(SemanticEntry.project_id == project.id)
        )
    ).scalar_one()
    assert entry.revision_number == 1
    message_count = await db_session.scalar(
        select(func.count()).select_from(Message).where(Message.conversation_id == conversation.id)
    )
    assert message_count == 1


@pytest.mark.asyncio
async def test_sheet_confirmation_rebuilds_working_copy_from_selected_sheet(
    client: AsyncClient,
    db_session: AsyncSession,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    project, conversation, run, source, question = await _seed_sheet_confirmation_run(
        client, db_session, tmp_path, monkeypatch
    )
    old_working_uri = str(source.working_uri)
    old_working = Path(old_working_uri)
    old_original_fingerprint = fingerprint_file(Path(str(source.source_uri)))
    assert pd.read_parquet(old_working).columns.tolist() == ["订单号", "门店ID", "商品", "实付"]

    response = await client.post(
        "/api/v1/chat/confirm",
        json={
            "analysis_run_id": str(run.id),
            "key": question["key"],
            "selected_option": "退款明细",
        },
    )

    assert response.status_code == 200, response.text
    await db_session.refresh(source)
    await db_session.refresh(run)
    assert source.working_uri != old_working_uri
    assert old_working.exists()
    selected = pd.read_parquet(Path(str(source.working_uri)))
    assert selected.columns.tolist() == ["退款单号", "渠道", "原因"]
    assert selected["退款单号"].tolist() == ["R-1", "R-2", "R-3"]
    assert fingerprint_file(Path(str(source.source_uri))) == old_original_fingerprint
    assert not any(
        item["key"].startswith("excel_sheet_selection")
        for item in source.profile_data["ambiguities"]
    )
    assert run.stage == "confirmation_received"
    assert run.checkpoint["confirmation_receipt"]["source_id"] == str(source.id)

    recipe = (
        (
            await db_session.execute(
                select(SanitationRecipeRecord)
                .where(SanitationRecipeRecord.data_source_id == source.id)
                .order_by(SanitationRecipeRecord.created_at.desc())
            )
        )
        .scalars()
        .first()
    )
    assert recipe is not None
    assert (
        next(item for item in recipe.operations if item["operation"] == "select_sheet")["sheet"]
        == "退款明细"
    )
    entry = (
        await db_session.execute(
            select(SemanticEntry).where(
                SemanticEntry.project_id == project.id,
                SemanticEntry.key == question["key"],
            )
        )
    ).scalar_one()
    assert entry.entry_type == "cleaning_rule"
    assert entry.value == "退款明细"
    assert entry.state == "confirmed"
    message = (
        await db_session.execute(select(Message).where(Message.conversation_id == conversation.id))
    ).scalar_one()
    assert message.content == "退款明细"


@pytest.mark.asyncio
async def test_sheet_confirmation_failure_preserves_trusted_copy_and_pending_question(
    client: AsyncClient,
    db_session: AsyncSession,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    project, conversation, run, source, question = await _seed_sheet_confirmation_run(
        client, db_session, tmp_path, monkeypatch
    )
    project_id = project.id
    conversation_id = conversation.id
    old_working_uri = str(source.working_uri)
    old_working_bytes = Path(old_working_uri).read_bytes()
    old_profile = copy.deepcopy(source.profile_data)
    recipe = (
        (
            await db_session.execute(
                select(SanitationRecipeRecord)
                .where(SanitationRecipeRecord.data_source_id == source.id)
                .order_by(SanitationRecipeRecord.created_at.desc())
            )
        )
        .scalars()
        .first()
    )
    assert recipe is not None
    old_recipe = copy.deepcopy(recipe.operations)

    def fail_after_creating_temp_file(source_path, output_dir, recipe_operations=None):
        del source_path, recipe_operations
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "incomplete.parquet").write_bytes(b"incomplete")
        raise ValueError("forced sheet rerun failure")

    monkeypatch.setattr("app.api.v1.projects.run_preflight", fail_after_creating_temp_file)
    response = await client.post(
        "/api/v1/chat/confirm",
        json={
            "analysis_run_id": str(run.id),
            "key": question["key"],
            "selected_option": "退款明细",
        },
    )

    assert response.status_code == 409, response.text
    await db_session.refresh(source)
    await db_session.refresh(run)
    await db_session.refresh(recipe)
    assert source.working_uri == old_working_uri
    assert Path(old_working_uri).read_bytes() == old_working_bytes
    assert source.profile_data == old_profile
    assert recipe.operations == old_recipe
    assert run.stage == "waiting_confirmation"
    assert any(item["key"] == question["key"] for item in source.profile_data["ambiguities"])
    entry_count = await db_session.scalar(
        select(func.count())
        .select_from(SemanticEntry)
        .where(
            SemanticEntry.project_id == project_id,
            SemanticEntry.key == question["key"],
        )
    )
    message_count = await db_session.scalar(
        select(func.count()).select_from(Message).where(Message.conversation_id == conversation_id)
    )
    assert entry_count == 0
    assert message_count == 0
    decisions_root = Path(old_working_uri).parent / "decisions"
    assert not decisions_root.exists() or not any(decisions_root.iterdir())


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("overrides", "expected_status"),
    [
        ({"selected_option": "自定义口径"}, 422),
        ({"key": "stale_confirmation_key"}, 409),
    ],
)
async def test_typed_confirmation_rejects_invalid_option_or_key_without_side_effects(
    client: AsyncClient,
    db_session: AsyncSession,
    overrides: dict[str, str],
    expected_status: int,
):
    project, conversation, run = await _seed_confirmation_run(db_session)

    response = await client.post(
        "/api/v1/chat/confirm",
        json=_command(run, **overrides),
    )

    assert response.status_code == expected_status
    await db_session.refresh(run)
    assert run.state == "waiting_confirmation"
    assert run.stage == "waiting_confirmation"
    knowledge_count = await db_session.scalar(
        select(func.count())
        .select_from(SemanticEntry)
        .where(SemanticEntry.project_id == project.id)
    )
    message_count = await db_session.scalar(
        select(func.count()).select_from(Message).where(Message.conversation_id == conversation.id)
    )
    assert knowledge_count == 0
    assert message_count == 0


@pytest.mark.asyncio
async def test_typed_confirmation_does_not_override_locked_definition(
    client: AsyncClient,
    db_session: AsyncSession,
):
    project, conversation, run = await _seed_confirmation_run(db_session)
    locked = SemanticEntry(
        project_id=project.id,
        key="revenue_refund_policy",
        value="保留退款订单",
        entry_type="business_rule",
        state="locked",
        confidence=1,
        evidence=[{"kind": "user_lock"}],
        source="user",
    )
    db_session.add(locked)
    await db_session.commit()

    response = await client.post("/api/v1/chat/confirm", json=_command(run))

    assert response.status_code == 409
    await db_session.refresh(locked)
    await db_session.refresh(run)
    assert locked.value == "保留退款订单"
    assert locked.state == "locked"
    assert run.stage == "waiting_confirmation"
    message_count = await db_session.scalar(
        select(func.count()).select_from(Message).where(Message.conversation_id == conversation.id)
    )
    assert message_count == 0


@pytest.mark.asyncio
async def test_typed_confirmation_reuses_one_legacy_alias_row_and_migrates_its_key(
    client: AsyncClient,
    db_session: AsyncSession,
):
    project, _conversation, run = await _seed_confirmation_run(db_session)
    legacy = SemanticEntry(
        project_id=project.id,
        key="refund_policy",
        value="保留退款订单",
        entry_type="business_rule",
        state="confirmed",
        confidence=1,
        source="user",
    )
    db_session.add(legacy)
    await db_session.commit()

    response = await client.post(
        "/api/v1/chat/confirm",
        json=_command(run, key="refund_handling"),
    )

    assert response.status_code == 200, response.text
    await db_session.refresh(legacy)
    await db_session.refresh(run)
    assert legacy.key == "revenue_refund_policy"
    assert legacy.value == "扣除退款"
    assert run.checkpoint["confirmation_receipt"]["key"] == "revenue_refund_policy"
    entry_count = await db_session.scalar(
        select(func.count())
        .select_from(SemanticEntry)
        .where(SemanticEntry.project_id == project.id)
    )
    assert entry_count == 1


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("is_active", "validity"),
    [(False, "active"), (True, "stale")],
)
async def test_typed_confirmation_reactivates_old_locked_row_without_reusing_its_proof(
    client: AsyncClient,
    db_session: AsyncSession,
    is_active: bool,
    validity: str,
):
    project, _conversation, run = await _seed_confirmation_run(db_session)
    old = SemanticEntry(
        project_id=project.id,
        key="refund_policy",
        value="保留退款订单",
        entry_type="business_rule",
        state="locked",
        confidence=1,
        definition={
            "version": 1,
            "kind": "business_rule_strategy",
            "rule_key": "refund_policy",
            "selected_option": "保留退款订单",
            "action": {
                "kind": "identity",
                "column": "退款状态",
                "observed_values": ["否", "已退款"],
            },
        },
        validity=validity,
        execution_state="verified",
        execution_details={"status": "verified", "proof": "old"},
        source="user",
        is_active=is_active,
    )
    db_session.add(old)
    await db_session.commit()

    response = await client.post(
        "/api/v1/chat/confirm",
        json=_command(run, key="refund_handling"),
    )

    assert response.status_code == 200, response.text
    await db_session.refresh(old)
    assert old.key == "revenue_refund_policy"
    assert old.value == "扣除退款"
    assert old.state == "confirmed"
    assert old.is_active is True
    assert old.validity == "unverified"
    assert old.definition is None
    assert old.execution_state == "definition_only"
    assert old.execution_details["status"] == "definition_only"
    assert old.revision_number == 1


@pytest.mark.asyncio
async def test_typed_confirmation_rejects_conflicting_rows_for_the_same_decision_slot(
    client: AsyncClient,
    db_session: AsyncSession,
):
    project, conversation, run = await _seed_confirmation_run(db_session)
    report = dict(run.report)
    report["confirmation"] = {
        **report["confirmation"],
        "key": "refund_handling",
    }
    run.report = report
    canonical = SemanticEntry(
        project_id=project.id,
        key="revenue_refund_policy",
        value="扣除退款",
        entry_type="business_rule",
        state="confirmed",
        confidence=1,
        source="user",
    )
    alias = SemanticEntry(
        project_id=project.id,
        key="refund_policy",
        value="保留退款订单",
        entry_type="business_rule",
        state="confirmed",
        confidence=1,
        source="user",
    )
    db_session.add_all([canonical, alias])
    await db_session.commit()

    response = await client.post(
        "/api/v1/chat/confirm",
        json=_command(run, key="refund_policy"),
    )

    assert response.status_code == 409, response.text
    await db_session.refresh(run)
    assert run.stage == "waiting_confirmation"
    assert run.checkpoint.get("confirmation_receipt") is None
    values = {
        entry.key: entry.value
        for entry in (
            await db_session.execute(
                select(SemanticEntry).where(SemanticEntry.project_id == project.id)
            )
        ).scalars()
    }
    assert values == {
        "refund_policy": "保留退款订单",
        "revenue_refund_policy": "扣除退款",
    }
    message_count = await db_session.scalar(
        select(func.count()).select_from(Message).where(Message.conversation_id == conversation.id)
    )
    assert message_count == 0


@pytest.mark.asyncio
async def test_confirmation_continues_and_completes_the_same_run(
    client: AsyncClient,
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
):
    project, conversation, run = await _seed_confirmation_run(db_session)
    response = await client.post("/api/v1/chat/confirm", json=_command(run))
    assert response.status_code == 200, response.text

    service = ExecutionService(db_session, project_id=project.id)

    class Inputs:
        model_config = {"model": "test"}
        history = [
            {
                "role": "assistant",
                "content": "计算收入时，退款订单需要扣除吗？",
                "confirmation": run.report["confirmation"],
            },
            {"role": "user", "content": "扣除退款"},
        ]

    class FakeEngine:
        async def execute(self, *, query, history, stop_checker):
            del history, stop_checker
            assert query.startswith(run.query)
            assert "扣除退款" in query
            yield SSEEvent.progress("investigating", "按确认口径继续")
            yield SSEEvent.result(
                "调查完成",
                data=[{"category": "A类", "store": "一店", "orders": 12}],
                report={"status": "completed", "title": "门店调查", "summary": "已完成"},
                analysis_state="completed",
                tool_history=[
                    {
                        "kind": "business_rule_application",
                        "rule_key": "revenue_refund_policy",
                    },
                    {
                        "kind": "file_sql",
                        "sql": "SELECT category, store, COUNT(*) AS orders FROM orders GROUP BY 1, 2",
                        "result_name": "store_order_summary",
                    },
                    {
                        "kind": "validation",
                        "result_name": "store_order_summary",
                        "profile": {
                            "columns": ["category", "store", "orders"],
                            "keys": {"category": {}, "store": {}},
                            "numeric": {"orders": {"count": 1}},
                            "truncated": False,
                        },
                    },
                ],
                knowledge_proposals=[],
            )

    async def fake_load_inputs(**kwargs):
        del kwargs
        return Inputs()

    async def fake_build_engine(inputs, *, run: AnalysisRun, resume_checkpoint):
        del inputs
        assert str(run.id) == response.json()["data"]["resume_run_id"]
        assert resume_checkpoint is None
        return FakeEngine()

    monkeypatch.setattr(service, "_load_execution_inputs", fake_load_inputs)
    monkeypatch.setattr(service, "_build_engine", fake_build_engine)

    events = [
        event
        async for event in service.execute_stream(
            query="不应创建新调查",
            conversation_id=conversation.id,
            resume_run_id=run.id,
        )
    ]

    result_event = next(event for event in events if event.type.value == "result")
    assert result_event.data["confirmed_corrections"][0]["key"] == "revenue_refund_policy"
    assert result_event.data["confirmed_corrections"][0]["task_query"] == run.query
    runs = list(
        (
            await db_session.execute(
                select(AnalysisRun).where(AnalysisRun.project_id == project.id)
            )
        ).scalars()
    )
    assert [item.id for item in runs] == [run.id]
    await db_session.refresh(run)
    assert run.state == "completed"
    assert run.checkpoint["confirmation_receipt_status"] == "consumed"
    assert run.checkpoint["confirmation_receipt"]["selected_value"] == "扣除退款"
    await db_session.refresh(project)
    assert project.extra_data["golden_scenarios"][0]["query"] == run.query
    verified = (
        await db_session.execute(
            select(SemanticEntry).where(
                SemanticEntry.project_id == project.id,
                SemanticEntry.entry_type == "verified_query",
            )
        )
    ).scalar_one()
    assert verified.state == "confirmed"


@pytest.mark.asyncio
@pytest.mark.parametrize("evidence_key", ["revenue_refund_policy", "refund_policy"])
async def test_executable_confirmation_accepts_canonical_or_alias_application_evidence(
    db_session: AsyncSession,
    evidence_key: str,
):
    project = Project(name="确认执行证据兼容")
    conversation = Conversation(title="收入调查", status="active")
    db_session.add_all([project, conversation])
    await db_session.flush()
    definition = {
        "version": 1,
        "kind": "business_rule_strategy",
        "rule_key": "revenue_refund_policy",
        "selected_option": "扣除退款",
        "action": {
            "kind": "value_filter",
            "column": "退款状态",
            "operator": "exclude",
            "values": ["已退款"],
            "observed_values": ["否", "已退款"],
        },
    }
    entry = SemanticEntry(
        project_id=project.id,
        key="revenue_refund_policy",
        value="扣除退款",
        entry_type="business_rule",
        state="confirmed",
        confidence=1,
        definition=definition,
        validity="active",
        execution_state="needs_validation",
        source="user",
    )
    run = AnalysisRun(
        project_id=project.id,
        conversation_id=conversation.id,
        query="收入是多少",
        state="understanding",
        stage="investigating",
        checkpoint={
            "confirmation_receipt": {
                "key": "revenue_refund_policy",
                "selected_value": "扣除退款",
                "applied": True,
                "conflict": False,
            },
            "confirmation_receipt_status": "in_progress",
            "continuation_kind": "confirmation",
        },
    )
    db_session.add_all([entry, run])
    await db_session.flush()
    await append_semantic_revision(
        db_session,
        entry,
        mutation_kind="created",
        actor_source="user",
        reason="建立待验证业务口径",
    )
    run.checkpoint = {
        **run.checkpoint,
        "confirmation_receipt": {
            **run.checkpoint["confirmation_receipt"],
            "value": entry.value,
            "semantic_entry_id": str(entry.id),
            "active_revision_id": str(entry.active_revision_id),
            "definition_hash": stable_payload_hash(entry.definition),
            "value_hash": stable_payload_hash(entry.value),
        },
    }
    await db_session.commit()
    definition_hash = stable_payload_hash(definition)
    service = ExecutionService(db_session, project_id=project.id)

    outcome = await service._persist_project_result(
        run,
        {
            "analysis_state": "completed",
            "result_name": "filtered_orders",
            "data": [{"订单号": "O-1", "退款状态": "否"}],
            "report": {
                "status": "completed",
                "title": "收入",
                "summary": "已按退款口径计算。",
            },
            "tool_history": [
                {
                    "kind": "business_rule_application",
                    "rule_key": evidence_key,
                    "rule_value": "扣除退款",
                    "semantic_entry_id": str(entry.id),
                    "active_revision_id": str(entry.active_revision_id),
                    "definition_hash": definition_hash,
                    "source_result": "orders",
                    "result_name": "filtered_orders",
                    "action_kind": "value_filter",
                    "before_rows": 2,
                    "after_rows": 1,
                    "excluded_rows": 1,
                    "input_hash": "a" * 64,
                    "output_hash": "b" * 64,
                },
                {
                    "kind": "validation",
                    "result_name": "filtered_orders",
                    "result_hash": "b" * 64,
                    "profile": {
                        "columns": ["订单号", "退款状态"],
                        "keys": {"订单号": {}},
                        "numeric": {},
                        "truncated": False,
                    },
                },
            ],
            "confirmed_corrections": [
                {
                    "key": "revenue_refund_policy",
                    "selected_value": "扣除退款",
                    "applied": True,
                    "conflict": False,
                    "task_query": run.query,
                }
            ],
            "knowledge_proposals": [],
        },
    )

    assert outcome.accepted is True
    await db_session.refresh(run)
    await db_session.refresh(entry)
    assert run.state == "completed"
    application = next(
        item
        for item in run.checkpoint["tool_history"]
        if item["kind"] == "business_rule_application"
    )
    assert application["rule_key"] == "revenue_refund_policy"
    assert entry.execution_state == "verified"
    assert len(await discover_report_correction_targets(db_session, run)) == 1


@pytest.mark.asyncio
async def test_definition_only_confirmation_does_not_require_application_evidence(
    db_session: AsyncSession,
):
    project = Project(name="仅定义确认")
    conversation = Conversation(title="业务定义", status="active")
    db_session.add_all([project, conversation])
    await db_session.flush()
    entry = SemanticEntry(
        project_id=project.id,
        key="revenue_refund_policy",
        value="扣除退款",
        entry_type="business_rule",
        state="confirmed",
        confidence=1,
        definition=None,
        validity="unverified",
        execution_state="definition_only",
        source="user",
    )
    run = AnalysisRun(
        project_id=project.id,
        conversation_id=conversation.id,
        query="记住退款口径",
        state="understanding",
        stage="investigating",
    )
    db_session.add_all([entry, run])
    await db_session.commit()
    service = ExecutionService(db_session, project_id=project.id)

    outcome = await service._persist_project_result(
        run,
        {
            "analysis_state": "completed",
            "report": {
                "status": "completed",
                "title": "口径已记录",
                "summary": "本次仅记录业务定义。",
            },
            "tool_history": [],
            "confirmed_corrections": [
                {
                    "key": "refund_handling",
                    "selected_value": "扣除退款",
                    "applied": True,
                    "conflict": False,
                }
            ],
            "knowledge_proposals": [],
        },
    )

    assert outcome.accepted is True
    await db_session.refresh(run)
    assert run.state == "completed"
    assert run.report["status"] == "completed"


@pytest.mark.asyncio
async def test_confirmation_continuation_can_retry_before_first_safe_result(
    client: AsyncClient,
    db_session: AsyncSession,
):
    project, conversation, run = await _seed_confirmation_run(db_session)
    response = await client.post("/api/v1/chat/confirm", json=_command(run))
    assert response.status_code == 200, response.text
    service = ExecutionService(db_session, project_id=project.id)

    prepared, query, checkpoint, receipt = await service._prepare_analysis_run(
        query="继续",
        conversation_id=conversation.id,
        resume_run_id=run.id,
    )
    assert prepared is run
    assert "扣除退款" in query
    assert checkpoint is None
    assert receipt and receipt["selected_value"] == "扣除退款"
    assert run.checkpoint["continuation_kind"] == "confirmation"
    assert run.checkpoint["resumable"] is True

    await service._mark_run_needs_attention(run, RuntimeError("模型连接暂时失败"))
    await db_session.refresh(run)
    assert run.state == "needs_attention"
    assert run.checkpoint["continuation_kind"] == "confirmation"
    assert run.checkpoint["resumable"] is True

    (
        retried,
        retried_query,
        retried_checkpoint,
        retried_receipt,
    ) = await service._prepare_analysis_run(
        query="再试一次",
        conversation_id=conversation.id,
        resume_run_id=run.id,
    )
    assert retried is run
    assert "扣除退款" in retried_query
    assert retried_checkpoint is None
    assert retried_receipt and retried_receipt["selected_value"] == "扣除退款"


@pytest.mark.asyncio
async def test_safe_checkpoint_resume_keeps_typed_confirmation_context(
    db_session: AsyncSession,
):
    project, conversation, run = await _seed_confirmation_run(db_session)
    receipt = {
        "key": "refund_handling",
        "selected_value": "扣除退款",
        "applied": True,
        "conflict": False,
        "task_query": run.query,
    }
    run.state = "needs_attention"
    run.stage = "paused"
    run.checkpoint = {
        "version": 1,
        "revision": 2,
        "resumable": True,
        "confirmation_receipt": receipt,
        "confirmation_receipt_status": "consumed",
    }
    await db_session.commit()
    service = ExecutionService(db_session, project_id=project.id)

    prepared, effective_query, checkpoint, restored_receipt = await service._prepare_analysis_run(
        query="继续",
        conversation_id=conversation.id,
        resume_run_id=run.id,
    )

    assert prepared is run
    assert "revenue_refund_policy = 扣除退款" in effective_query
    assert checkpoint and checkpoint["revision"] == 2
    assert restored_receipt == {**receipt, "key": "revenue_refund_policy"}
    assert run.checkpoint["confirmation_receipt"]["key"] == "revenue_refund_policy"
